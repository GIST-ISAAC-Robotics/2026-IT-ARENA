from pathlib import Path
import json, re, zipfile, xml.etree.ElementTree as ET
import openpyxl

root=Path(__file__).resolve().parents[2]
out=Path(__file__).parent
book=out/'IT_ARENA_센서_예산안_2026-09-05_ToF수정.xlsx'
# Artifact Tool has no documented worksheet reorder API. Change only workbook
# sheet-node order, leaving all authored cell data, styles and formulas intact.
ns={'m':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
order=['구성 비교','가격순 비교','스테레오 비교','단가와 출처','차체 견적 대조','예산 조건']
with zipfile.ZipFile(book) as z:
    files={item.filename:(item,z.read(item.filename)) for item in z.infolist()}
raw=files['xl/workbook.xml'][1].decode('utf-8')
match=re.search(r'(<x:sheets\b[^>]*>)(.*?)(</x:sheets>)',raw,re.S)
assert match, 'workbook sheets missing'
nodes=re.findall(r'<x:sheet\b[^>]*/>',match.group(2))
byname={re.search(r'name="([^"]+)"',n).group(1):n for n in nodes}
assert set(byname)==set(order)
raw=raw[:match.start(2)]+''.join(byname[n] for n in order)+raw[match.end(2):]
temp=book.with_suffix('.reordered.xlsx')
with zipfile.ZipFile(temp,'w',zipfile.ZIP_DEFLATED) as z:
    for name,(item,data) in files.items():
        z.writestr(item,raw.encode('utf-8') if name=='xl/workbook.xml' else data)
temp.replace(book)
with zipfile.ZipFile(book) as z:
    assert all(z.read(name)==data for name,(_,data) in files.items() if name!='xl/workbook.xml')
data=openpyxl.load_workbook(book,data_only=True)
formula=openpyxl.load_workbook(book,data_only=False)
assert data.sheetnames==order
expected=json.loads((out/'results.json').read_text(encoding='utf-8'))
actual=[[data['구성 비교'].cell(r,c).value for c in range(1,8)] for r in range(5,5+len(expected))]
assert actual==expected
sorted_actual=[[data['가격순 비교'].cell(r,c).value for c in range(1,8)] for r in range(5,5+len(expected))]
assert sorted_actual==sorted(expected,key=lambda r:r[3]), 'SORT export cache mismatch'
sort_formula=formula['가격순 비교']['A5'].value
sort_text=sort_formula.text if hasattr(sort_formula,'text') else sort_formula
assert 'SORT(' in sort_text
assert data['차체 견적 대조']['D22'].value==221150
assert data['차체 견적 대조']['D24'].value==268550
assert data['스테레오 비교']['B7'].value==605000
print('Saved file verified:',data.sheetnames,'20 sorted rows, SORT formula, chassis totals')


assert len(expected)==20
assert {row[0]:row[2] for row in actual if row[0] in ['L2-0','H-0','Hfast-0']} == {'L2-0':0,'H-0':0,'Hfast-0':0}
assert next(row[3] for row in actual if row[0]=='H-0')==696000
assert next(row[3] for row in actual if row[0]=='Hfast-0')==1336200
for row in data['스테레오 비교'].iter_rows(min_row=5,max_row=14,min_col=1,max_col=9):
    assert all(c.value is not None for c in row)
print('10 camera mode/range/source records and 3 ToF-free hybrids verified.')

assert all(row[2] == (0 if 'C1' in row[1] or 'S3 20Hz' in row[1] else 4) for row in actual)
assert next(row[3] for row in actual if row[0]=='RV-4')==204160
assert next(row[3] for row in actual if row[0]=='RA-4')==237160
print('All LiDAR rows ToF0, both RGB-only rows ToF4 verified.')
